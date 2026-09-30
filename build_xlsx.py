import os
import json, glob, importlib.util, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
rows=json.load(open('merged.json')); scr=json.load(open('scr.json')); cls=json.load(open('cls.json'))
done={}
for f in sorted(glob.glob('dec*.py')):
    sp=importlib.util.spec_from_file_location(f[:-3],f); m=importlib.util.module_from_spec(sp); sp.loader.exec_module(m); done.update(m.D)
by={r['sym']:r for r in rows}
def offc(r):
    c=cls.get(r['sym']) or {}
    if c.get('basic'): return ' > '.join(x for x in [c.get('macro'),c.get('sector'),c.get('industry'),c.get('basic')] if x)
    s=scr.get(r['sym']) or {}
    return ' > '.join(s.get('cls') or []) or 'Not available'
def about(r):
    s=scr.get(r['sym']) or {}
    return (s.get('about') or '').replace('[1]','').replace('[2]','').strip()
wb=Workbook()
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F3A2E'); WR=Alignment(wrap_text=True,vertical='top')
def sheet(ws,head,data,widths):
    ws.append(head)
    for c in ws[1]: c.font=H; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center')
    for d in data: ws.append(d)
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=WR
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
conf={'H':'High','M':'Medium'}
W=[];R=[];K=[]
for s,(v,tgt,cf,why) in done.items():
    r=by[s]
    base=[r['co'],s,r['ind']]
    tail=[conf[cf],why,offc(r),r['tv_ind'] or 'Not on TradingView scanner',about(r),r['wl']]
    if v=='W': W.append(base+[tgt]+tail)
    elif v=='R': R.append(base+[tgt or '']+tail)
    else: K.append(base+[why,offc(r)])
key=lambda x:(x[2],x[0])
W.sort(key=lambda x:(x[2],0 if x[4]=='High' else 1,x[0])); R.sort(key=key); K.sort(key=key)
ws=wb.active; ws.title='Misclassified'
sheet(ws,['Company','TradingView symbol','Current industry','Correct industry','Confidence','Why','Official exchange classification (BSE/NSE)','TradingView industry','Business description (Screener)','Watchlist'],W,[34,20,30,30,11,55,48,26,70,34])
ws=wb.create_sheet('Needs your call')
sheet(ws,['Company','TradingView symbol','Current industry','Suggested industry (if any)','Confidence','Why it is unclear','Official exchange classification (BSE/NSE)','TradingView industry','Business description (Screener)','Watchlist'],R,[34,20,30,30,11,55,48,26,70,34])
ws=wb.create_sheet('Flagged but correct')
sheet(ws,['Company','TradingView symbol','Current industry','Why it stays','Official exchange classification (BSE/NSE)'],K,[34,20,30,50,60])
# not on scanner
NS=[[r['co'],r['sym'],r['ind'], (lambda v: {'W':'Misclassified','R':'Needs your call','K':'Correct'}.get(v[0],'') if v else 'Label fits')(done.get(r['sym']))] for r in rows if not r['tv_name']]
ws=wb.create_sheet('Not on TV scanner')
sheet(ws,['Company','TradingView symbol','Current industry','Audit result'],sorted(NS,key=lambda x:(x[2],x[0])),[40,22,34,18])
# summary
cnt=collections.Counter(r['ind'] for r in rows)
wc=collections.Counter(by[s]['ind'] for s,v in done.items() if v[0]=='W')
rc=collections.Counter(by[s]['ind'] for s,v in done.items() if v[0]=='R')
into=collections.Counter(v[1] for s,v in done.items() if v[0]=='W')
S=[[i,cnt[i],wc[i],rc[i],round(100*wc[i]/cnt[i],1),into[i]] for i in sorted(cnt)]
S.sort(key=lambda x:(-x[2],x[0]))
ws=wb.create_sheet('Summary by industry')
sheet(ws,['Industry','Names in watchlist','Misclassified (move out)','Needs your call','% misclassified','Names that should move in'],S,[42,12,14,12,12,14])
M=[["What was checked","All 4,707 names in the 6 'Industry - ...' TradingView watchlists (278 industries). The live watchlists were confirmed identical to Desktop\Industry-wise stocks.xlsx on 18 Sep 2026."],
["Independent evidence","(1) Official 4-level exchange classification from BSE for 4,196 BSE-listed names and from Screener.in (same NSE/BSE framework) for NSE-only names; (2) TradingView's own sector/industry; (3) the Screener business description."],
["Rule","For each of the 278 industries I defined which official 'basic industries' are acceptable. 1,167 names (25%) broke their industry's rule. Each of those was read and judged by hand against its business description."],
["Verdicts","Misclassified = the business clearly belongs in another of your 278 industries (the 'Correct industry' is always one of your existing buckets). Needs your call = shells, renamed companies whose official record is stale, or genuinely mixed businesses. Flagged but correct = the official label is coarse but your bucket fits."],
["Confidence","High = the description and official label both point clearly to the new bucket. Medium = the move is right but the target is a judgement between two of your buckets (e.g. Real Estate vs Construction EPC)."],
["Limit","3,540 names passed because their official label fits their bucket; those were not read one by one. Errors inside a coarse official label (e.g. an API maker filed under Pharma Formulations, both labelled 'Pharmaceuticals') are therefore not caught."],
["Not changed","Nothing in TradingView or in your workbooks was modified."]]
ws=wb.create_sheet('Method')
for m in M: ws.append(m)
ws.column_dimensions['A'].width=22; ws.column_dimensions['B'].width=120
for row in ws.iter_rows():
    for c in row: c.alignment=WR
    row[0].font=Font(bold=True)
out=os.environ.get('AUDIT_XLSX', 'Industry classification audit - 18 Sep 2026.xlsx')
wb.save(out)
print('saved',out,'W',len(W),'R',len(R),'K',len(K),'notscan',len(NS))
print('W high',sum(1 for x in W if x[4]=='High'),'medium',sum(1 for x in W if x[4]=='Medium'))
print(S[:25])
